"""АВ cabinet on the real rules (config/kv_rules.json) and synthetic agents: scales, bounds, newcomers,
Доп.АВ1–3, the КАСКО cap, unrated lines, planner, effects, scenarios, versions, inputs, API."""

import json
import os
import shutil
from datetime import datetime

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from metrixs_ml import settings
from metrixs_ml.api import create_app
from metrixs_ml.kv.engine import Additions, KvDataError, KvInputs, evaluate, income, is_newcomer, parse_quarter
from metrixs_ml.kv.inputs import build_kv_inputs
from metrixs_ml.kv.rules import RulesCache, RulesError
from metrixs_ml.kv.service import KvService

IFL = "ИМУЩЕСТВО ФИЗ. ЛИЦ"
Q2, Q3, Q4 = parse_quarter("2026Q2"), parse_quarter("2026Q3"), parse_quarter("2026Q4")
OLD = {"branch": "ФИЛИАЛ Т", "group": "2 группа", "firstAd": "2015-01-01"}


@pytest.fixture(scope="module")
def rule():
    book, _ = RulesCache(settings.KV_RULES_FILE).get()
    return book.for_quarter(Q3)


def _q(kind: str, kasko: float, ifl: float = 0.0, due: int = 0, prolonged: int | None = None, estimated: bool = False,
       **other: float) -> dict:
    fact = kind == "fact"
    lines = {"КАСКО": kasko, IFL: ifl, **{k.replace("_", " "): v for k, v in other.items()}}
    return {
        "kind": kind,
        "lines": {k: {"toDate": v if fact else 0.0, "remainder": 0.0 if fact else v, "contracts": 10} for k, v in lines.items()},
        "prolongDue": due,
        "prolonged": prolonged,
        "prolongDueEstimated": estimated,
    }


def _agent(kasko: float = 400_000, ifl: float = 200_000, due: int = 10, prolonged: int = 8) -> dict:
    return {
        "2026Q1": _q("fact", kasko, ifl),
        "2026Q2": _q("fact", kasko, ifl, due=due, prolonged=prolonged),
        "2026Q3": _q("forecast", kasko, ifl, due=due, estimated=True),
        "2026Q4": _q("forecast", kasko, ifl),
    }


def _inputs(quarters: dict, info: dict | None = None) -> dict:
    return {
        "source": "test.xlsx",
        "asOf": "2026-06-30",
        "currentQuarter": "2026Q3",
        "currentQuarterHasData": False,
        "elapsedShare": 0.0,
        "rowsAfterAsOfExcluded": 0,
        "prolongationOutcomesAvailable": True,
        "agents": sorted(quarters),
        "lines": ["КАСКО", IFL],
        "agentInfo": {a: (info or {}).get(a, OLD) for a in quarters},
        "quarters": quarters,
        "forecastQuality": {"scores": [], "chosen": {"method": "mean3", "title": "среднее", "wape": 0.3}, "months": 6,
                            "note": "тест"},
    }


def _service(tmp_path, quarters: dict, info: dict | None = None, rules_path=None) -> KvService:
    path = rules_path or tmp_path / "kv_rules.json"
    if rules_path is None:
        shutil.copy(settings.KV_RULES_FILE, path)
    return KvService(KvInputs(_inputs(quarters, info)), RulesCache(path))


def _eval(rule, reporting: dict, previous: dict | None = None, info: dict = OLD, quarter=Q3, additions=None) -> dict:
    return evaluate(rule, info, quarter, reporting, previous or _q("fact", 0), additions=additions)


def _rates(result: dict) -> dict:
    return {r["line"]: r for r in result["rates"]}


# --- base rate --------------------------------------------------------------------------------------

@pytest.mark.parametrize(("group", "dvs", "expected"), [
    ("2 группа", 299_999, 0.23), ("2 группа", 300_000, 0.24), ("2 группа", 800_000, 0.26),
    ("1 группа", 99_999, 0.22), ("3 группа", 300_000, 0.23), ("3 группа", 1_100_000, 0.26),
])
def test_kasko_scale_depends_on_group_and_bound_is_exclusive(rule, group, dvs, expected):
    result = _eval(rule, _q("fact", dvs), info={**OLD, "group": group})

    assert result["components"]["kaskoBase"]["rate"] == expected


def test_newcomer_gets_25_percent_until_the_right_quarter(rule):
    assert is_newcomer("2026-04-01", Q2) and not is_newcomer("2026-04-01", Q3)  # 3 full months left in Q2
    assert is_newcomer("2026-05-01", Q2) and not is_newcomer("2026-05-01", Q3)  # exactly 2 full months
    assert is_newcomer("2026-05-15", Q3) and not is_newcomer("2026-05-15", Q4)  # < 2 months: until the next quarter end
    result = _eval(rule, _q("fact", 1_000), info={**OLD, "firstAd": "2026-05-15"})
    assert result["components"]["kaskoBase"]["rate"] == 0.25 and result["components"]["kaskoBase"]["newcomer"]


def test_fixed_base_rates_and_unrated_lines(rule):
    result = _eval(rule, _q("fact", 0, 100_000, СПЕЦВИДЫ=50_000, НОВЫЙ_ВИД=10_000))
    rates = _rates(result)

    assert rates[IFL]["base"] == 0.35 and rates["КИС"]["base"] == 0.5
    assert rates["СПЕЦВИДЫ"]["total"] is None and rates["НОВЫЙ ВИД"]["total"] is None
    assert result["metrics"]["dvs"] == 160_000 and result["metrics"]["priority"] == 160_000  # unrated still count in ДВС
    inc = income(rule, result["rates"], result["amountsReporting"])
    assert {u["line"] for u in inc["unrated"]} == {"СПЕЦВИДЫ", "НОВЫЙ ВИД"}
    assert inc["kv"] == pytest.approx(100_000 * rates[IFL]["total"])


# --- Доп.АВ1–3 and the cap ------------------------------------------------------------------------------

def test_dynamics_uses_reporting_quarter_bands_and_100k_floor(rule):
    # R = 2026Q2 → bands 35 / 50 / 60 %; previous ДВС 50 000 → base 100 000
    result = _eval(rule, _q("fact", 150_000), _q("fact", 50_000))
    m = result["metrics"]

    assert m["dynamicsBase"] == 100_000 and m["dynamicsFloorApplied"]
    assert m["dynamics"] == pytest.approx(0.5) and result["components"]["d1"]["value"] == 0.02  # 50 % is not < 50 %
    assert _eval(rule, _q("fact", 149_999), _q("fact", 50_000))["components"]["d1"]["value"] == 0.01
    assert _eval(rule, _q("fact", 0), _q("fact", 0))["metrics"]["dynamics"] == 0


def test_penetration_must_be_strictly_above_threshold(rule):
    at = _eval(rule, _q("fact", 901_000, 99_000))  # exactly 9.9 %
    above = _eval(rule, _q("fact", 900_999, 99_001))

    assert at["components"]["d2"]["value"] == 0 and above["components"]["d2"]["value"] == 0.02


def test_prolongation_threshold_and_missing_result(rule):
    def d3(due, prolonged):
        return _eval(rule, _q("fact", 100_000, due=due, prolonged=prolonged))["components"]["d3"]["value"]

    assert d3(10_000, 6_999) == 0.02 and d3(10_000, 6_998) == 0 and d3(10, 7) == 0.02
    missing = _eval(rule, _q("forecast", 100_000, due=10, estimated=True))["metrics"]["prolongation"]
    assert missing["rate"] is None and missing["missing"] and missing["dueEstimated"]


def test_kasko_total_is_capped_at_30_percent(rule):
    result = _eval(rule, _q("fact", 2_000_000, 500_000, due=10, prolonged=10), _q("fact", 1_000_000))
    rates = _rates(result)

    assert rates["КАСКО"]["base"] + rates["КАСКО"]["d1"] + rates["КАСКО"]["d2"] + rates["КАСКО"]["d3"] == pytest.approx(0.33)
    assert rates["КАСКО"]["total"] == 0.3 and rates["КАСКО"]["capped"]
    assert rates[IFL]["total"] == pytest.approx(0.42) and not rates[IFL]["capped"]


# --- planner and effects ------------------------------------------------------------------------------

def test_priority_sale_moves_numerator_and_denominator(tmp_path):
    svc = _service(tmp_path, {"A": _agent(kasko=950_000, ifl=50_000)})
    step = next(s for s in svc.cabinet("A", "2026Q3")["plan"]["steps"] if s["component"] == "d2")
    v = step["needs"]["priority"]

    # (50 000 + v) / (1 000 000 + v) > 9.9 %, not (50 000 + v) / 1 000 000
    assert 49_000 < v < 55_000
    assert svc.rate("A", Q4, additions=Additions(priority=v))["components"]["d2"]["value"] == 0.02
    assert svc.rate("A", Q4, additions=Additions(priority=v - 2))["components"]["d2"]["value"] == 0


def test_kasko_sale_raises_volume_but_lowers_penetration(tmp_path):
    svc = _service(tmp_path, {"A": _agent(kasko=900_000, ifl=100_001)})
    before, after = svc.rate("A", Q4), svc.rate("A", Q4, additions=Additions(kasko=100_000))

    assert after["metrics"]["dvs"] > before["metrics"]["dvs"]
    assert before["components"]["d2"]["value"] == 0.02 and after["components"]["d2"]["value"] == 0


def test_plan_options_are_verified_and_effects_split(tmp_path):
    svc = _service(tmp_path, {"A": _agent()})
    plan = svc.cabinet("A", "2026Q3")["plan"]

    assert plan["available"] and plan["salesQuarter"] == "2026Q3"
    targets = {s["component"] for s in plan["steps"]}
    for option in plan["options"]:
        after = svc.rate("A", Q4, additions=Additions(kasko=option["kasko"], priority=option["priority"],
                                                      prolonged=option["prolonged"]))
        for step in plan["steps"]:
            key = "rate" if step["component"] == "kaskoBase" else "value"
            assert after["components"][step["component"]][key] >= step["to"] - 1e-9, step["component"]
        e = option["effects"]
        assert e["total"] == pytest.approx(e["salesEffect"] + e["rateEffect"])
    assert {"kaskoBase", "d1", "d3"} <= targets


def test_sales_effect_uses_open_quarter_rates_and_rate_effect_the_next_base(tmp_path):
    svc = _service(tmp_path, {"A": _agent()})
    adds = Additions(priority=100_000)

    sim = svc.simulate("A", "2026Q3", adds)
    ifl_q3 = _rates(svc.rate("A", Q3))[IFL]["total"]
    before, after = _rates(svc.rate("A", Q4)), _rates(svc.rate("A", Q4, additions=adds))
    expected_rate_effect = 400_000 * (after["КАСКО"]["total"] - before["КАСКО"]["total"]) + 200_000 * (
        after[IFL]["total"] - before[IFL]["total"])

    assert sim["effects"]["salesEffect"] == pytest.approx(100_000 * ifl_q3)
    assert sim["effects"]["rateEffect"] == pytest.approx(expected_rate_effect)
    assert sim["baseline"]["incomeNext"]["kv"] == pytest.approx(svc.cabinet("A", "2026Q3")["income"]["next"]["base"]["kv"])


def test_prolongation_plan_counts_contracts_due(tmp_path):
    svc = _service(tmp_path, {"A": _agent(due=27)})
    step = next(s for s in svc.cabinet("A", "2026Q3")["plan"]["steps"] if s["component"] == "d3")

    assert step["needs"] == {"prolonged": 19, "due": 27, "dueEstimated": True}  # ceil(0.6999 × 27)
    assert svc.rate("A", Q4, additions=Additions(prolonged=19))["components"]["d3"]["value"] == 0.02


# --- periods, scenarios, data guards --------------------------------------------------------------------

def test_fact_and_forecast_status_and_scenarios(tmp_path):
    svc = _service(tmp_path, {"A": _agent()})
    cab = svc.cabinet("A", "2026Q3")

    assert cab["currentRate"]["status"] == "fact" and cab["nextRate"]["status"] == "forecast"
    cur = cab["income"]["current"]
    assert cur["cautious"]["base"] == pytest.approx(0.9 * cur["base"]["base"])
    assert cur["optimistic"]["base"] == pytest.approx(1.1 * cur["base"]["base"])
    assert svc.rate("A", Q3, "cautious")["rates"] == svc.rate("A", Q3)["rates"]  # both basis quarters are facts


def test_closed_quarters_and_unknown_inputs_are_rejected(tmp_path):
    svc = _service(tmp_path, {"A": _agent()})

    with pytest.raises(KvDataError):
        svc.rate("A", Q3, additions=Additions(priority=1))  # R = 2026Q2 is closed
    with pytest.raises(KvDataError):
        svc.cabinet("Чужой", "2026Q3")
    with pytest.raises(KvDataError):
        svc.cabinet("A", "2026Q4")
    assert svc.options()["quarters"] == ["2026Q3"]


def test_excluded_branch_is_rejected(tmp_path):
    svc = _service(tmp_path, {"A": _agent()}, info={"A": {**OLD, "branch": "ФИЛИАЛ В Г. САНКТ-ПЕТЕРБУРГ"}})

    with pytest.raises(KvDataError, match="не применяется"):
        svc.cabinet("A", "2026Q3")


# --- rule versions -------------------------------------------------------------------------------------

def test_rule_versions_switch_gap_and_reload(tmp_path):
    path = tmp_path / "kv_rules.json"
    book = json.loads(settings.KV_RULES_FILE.read_text(encoding="utf-8"))
    old = {**book["versions"][0], "id": "old", "valid_to": "2025-12-31"}
    new = {**book["versions"][0], "id": "new", "valid_from": "2026-07-01", "lines": {
        **book["versions"][0]["lines"], IFL: {"title": "ИФЛ", "priority": True, "base_rate": 0.4}}}
    path.write_text(json.dumps({"versions": [old, new]}, ensure_ascii=False), encoding="utf-8")
    cache = RulesCache(path)

    assert cache.get()[0].for_quarter(Q3).id == "new"
    with pytest.raises(RulesError):
        cache.get()[0].for_quarter(Q2)  # 2026H1 not covered
    svc = KvService(KvInputs(_inputs({"A": _agent()})), cache)
    assert _rates(svc.cabinet("A", "2026Q3")["currentRate"])[IFL]["base"] == 0.4

    path.write_text("{broken", encoding="utf-8")
    os.utime(path, (1, 1))
    with pytest.raises(RulesError):
        cache.get()


# --- inputs from the export -----------------------------------------------------------------------------

def _export(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "сборы"
    ws.append(["Отч период", "Дата заключения", "Линия бизнеса", "Договор", "Филиал", "Агент", "Сумма сборов"])
    ends = [datetime(2025, 3, 31), datetime(2025, 6, 30), datetime(2025, 9, 30), datetime(2025, 12, 31),
            datetime(2026, 3, 31), datetime(2026, 6, 30), datetime(2026, 9, 30)]
    for i, end in enumerate(ends):
        ws.append([end, end, "КАСКО", f"K{i}", "ФИЛИАЛ Т", "Агент А", 100_000])
        ws.append([end, end, IFL, f"I{i}", "ФИЛИАЛ Т", "Агент А", 20_000])
    ws.append([datetime(2026, 9, 30), datetime(2026, 8, 1), "КАСКО", "K-open", "ФИЛИАЛ Т", "Агент А", 30_000])
    ws.append([datetime(2026, 9, 30), datetime(2026, 9, 20), "КАСКО", "K-late", "ФИЛИАЛ Т", "Агент А", 999_000])
    p = wb.create_sheet("пролонг")
    p.append([])  # header on the second row, as in the export
    p.append(["Агент", "Линия бизнеса", "Договор к пролонгации", "Дата окончания", "Пролонгировано всего"])
    for i in range(4):
        p.append(["Агент А", "КАСКО", f"P{i}", datetime(2026, 5, 10), 1 if i < 3 else 0])
    p.append(["Агент А", "КАСКО", "P-q3", datetime(2026, 8, 10), None])
    g = wb.create_sheet("группа филиала")
    g.append(["Филиал", "группа"])
    g.append(["ФИЛИАЛ Т", "Группа Канто"])
    a = wb.create_sheet("дата 1 АД агента")
    a.append(["Агент", "Дата 1 АД"])
    a.append(["Агент А", datetime(2015, 1, 1)])
    wb.save(path)


def test_inputs_use_only_records_dated_by_report_date(tmp_path):
    _export(tmp_path / "data.xlsx")
    data = build_kv_inputs(tmp_path / "data.xlsx", pd.Timestamp("2026-08-15"))
    quarters = data["quarters"]["Агент А"]

    assert data["currentQuarter"] == "2026Q3" and data["rowsAfterAsOfExcluded"] == 3  # K6, I6 and K-late
    assert data["agentInfo"]["Агент А"] == {"branch": "ФИЛИАЛ Т", "group": "Группа Канто", "firstAd": "2015-01-01"}
    assert quarters["2026Q2"]["kind"] == "fact" and quarters["2026Q2"]["lines"]["КАСКО"]["toDate"] == 100_000
    assert (quarters["2026Q2"]["prolongDue"], quarters["2026Q2"]["prolonged"]) == (4, 3)
    q3 = quarters["2026Q3"]
    assert q3["kind"] == "nowcast" and q3["lines"]["КАСКО"]["toDate"] == 30_000
    assert q3["lines"]["КАСКО"]["remainder"] == pytest.approx(100_000 * (1 - data["elapsedShare"]))  # flat history
    assert (q3["prolongDue"], q3["prolonged"]) == (1, None)  # due known, outcome missing
    assert quarters["2026Q4"]["kind"] == "forecast"


def test_data_group_is_mapped_onto_the_scale_by_a_labelled_assumption(rule):
    result = _eval(rule, _q("fact", 1_100_000), info={**OLD, "group": "Группа Канто"})

    assert result["components"]["kaskoBase"]["scaleGroup"] == "3 группа"
    assert result["components"]["kaskoBase"]["rate"] == 0.26
    assert rule.kasko.group_aliases_note


# --- API -----------------------------------------------------------------------------------------------

@pytest.fixture
def api(tmp_path):
    (tmp_path / "plans.json").write_text("[]", encoding="utf-8")
    (tmp_path / "metrics.json").write_text(json.dumps({"agents": ["A"], "scores": [], "chosen": "trend"}), encoding="utf-8")
    (tmp_path / "kv_inputs.json").write_text(json.dumps(_inputs({"A": _agent()}), ensure_ascii=False), encoding="utf-8")
    bonus = tmp_path / "bonus.json"
    bonus.write_text(json.dumps({"rate": 0.05, "threshold": 0, "cap_per_quarter": 1, "max_refund_share": 0.1, "refund_penalty": 0.5}))
    with TestClient(create_app(artifacts_dir=tmp_path, bonus_config_file=bonus, kv_rules_file=settings.KV_RULES_FILE)) as client:
        yield client


def test_api_cabinet_simulate_and_export(api):
    cab = api.get("/kv/cabinet", params={"agent": "A", "quarter": "2026Q3"}).json()["data"]
    assert cab["nextQuarter"] == "2026Q4" and cab["plan"]["available"]

    sim = api.post("/kv/simulate", json={"agent": "A", "quarter": "2026Q3", "extra_priority": 100_000, "prolonged": 7}).json()["data"]
    assert sim["scenario"]["nextRate"]["components"]["d3"] == 0.02 and sim["baseline"]["nextRate"]["components"]["d3"] == 0

    xlsx = api.get("/kv/export", params={"agent": "A", "quarter": "2026Q3", "extra_kasko": 50_000})
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    assert "attachment" in xlsx.headers["content-disposition"]


def test_api_rejects_bad_input(api):
    assert api.get("/kv/cabinet", params={"agent": "Чужой", "quarter": "2026Q3"}).status_code == 400
    assert api.post("/kv/simulate", json={"agent": "A", "quarter": "2026Q3", "extra_priority": -1}).status_code == 422
    assert api.post("/kv/simulate", json={"agent": "A", "quarter": "2026Q2"}).status_code == 400


# --- prolongation lines and main-page recommendations -------------------------------------------------

def test_prolongation_counts_only_configured_lines(rule):
    entry = {**_q("fact", 100_000), "prolongDue": 20, "prolonged": 10,
             "prolongDueByLine": {"КАСКО": 10, "КИС": 10}, "prolongedByLine": {"КАСКО": 7, "КИС": 3}}

    p = _eval(rule, entry)["metrics"]["prolongation"]

    assert rule.prolongation.lines == ["КАСКО"]
    assert (p["due"], p["prolonged"]) == (10, 7) and _eval(rule, entry)["components"]["d3"]["value"] == 0.02


def _with_portfolio(quarters: dict) -> dict:
    q3 = quarters["2026Q3"]
    quarters["2026Q3"] = {**q3, "prolongDue": 12, "prolongDueByLine": {"КАСКО": 10, "КИС": 2}}
    return quarters


def test_recommendations_start_from_the_kasko_portfolio(tmp_path):
    from metrixs_ml.kv.recommend import recommendations

    svc = _service(tmp_path, {"A": _with_portfolio(_agent())})
    svc.inputs.data["prolongContracts"] = {"A": [
        {"contract": f"K{i}", "line": "КАСКО", "product": "П1" if i < 6 else "П2", "endAt": f"2026-09-{10 + i}"}
        for i in range(10)
    ] + [{"contract": "M1", "line": "КИС", "product": "Ипотека", "endAt": "2026-09-01"}]}
    svc.inputs.data["productMix"] = {"A": {IFL: [{"product": "Платинум", "amount": 1.0, "contracts": 1}]}}

    r = recommendations(svc, "A")
    p, s = r["portfolio"], r["sell"]

    assert p["due"] == 10 and p["needed"] == 7  # only КАСКО, ceil(0.6999 × 10)
    assert [(g["product"], g["count"]) for g in p["groups"]] == [("П1", 6), ("П2", 4)]
    assert p["effects"]["rateEffect"] > 0 and p["effects"]["salesEffect"] == 0
    assert sum(row["planAmount"] for row in s["rows"]) == pytest.approx(s["total"], abs=0.05)
    assert s["rows"][-1]["line"] == "КАСКО" and s["rows"][0]["priority"]
    assert s["perRenewal"] == pytest.approx(s["total"] / 7)
    assert next(row for row in s["rows"] if row["line"] == IFL)["products"][0]["product"] == "Платинум"


def test_recommendations_count_contracts_behind_each_rate(tmp_path):
    from metrixs_ml.kv.recommend import recommendations

    svc = _service(tmp_path, {"A": _agent()})
    svc.inputs.data["prolongContracts"] = {"A": []}
    svc.inputs.data["productMix"] = {"A": {}}

    r = recommendations(svc, "A")
    planned = {row["line"]: row["planContracts"] or 0 for row in r["sell"]["rows"]}

    # 2026Q3's rate — 2026Q2 fact; 2026Q4's — 2026Q3 to date (10) + remainder ÷ average check (400 000 ÷ 40 000)
    assert r["rates"]
    for rate in r["rates"]:
        assert rate["contracts"] == {"current": 10, "next": 20, "nextWithPlan": 20 + planned.get(rate["line"], 0)}
