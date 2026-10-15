"""Excel export of an АВ cabinet calculation (and an optional «what if» scenario)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append(row)
    for i, _ in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 22


def _pct(v):
    return None if v is None else round(v * 100, 2)


def _rates_rows(rate: dict) -> list[list]:
    return [[rate["quarter"], r["title"], _pct(r["base"]), _pct(r["d1"]), _pct(r["d2"]), _pct(r["d3"]), _pct(r["total"]),
             "да" if r["capped"] else "", r["note"] or ""] for r in rate["rates"]]


def _metrics_rows(rate: dict) -> list[list]:
    m, c = rate["metrics"], rate["components"]
    p = m["prolongation"]
    return [
        [rate["quarter"], f"ДВС {rate['reporting']}", round(m["dvs"], 2), _pct(c["kaskoBase"]["rate"]), "база КАСКО"],
        [rate["quarter"], f"ДВС {rate['previous']}", round(m["dvsPrevious"], 2), None,
         "база динамики 100 000 ₽" if m["dynamicsFloorApplied"] else ""],
        [rate["quarter"], "Динамика ДВС, %", _pct(m["dynamics"]), _pct(c["d1"]["value"]), "Доп.АВ1"],
        [rate["quarter"], "Доля приоритетных, %", _pct(m["penetration"]), _pct(c["d2"]["value"]), "Доп.АВ2"],
        [rate["quarter"], f"Пролонгация {p['prolonged'] if p['prolonged'] is not None else '?'} из {p['due']}",
         _pct(p["rate"]), _pct(c["d3"]["value"]), p["missing"] or ("к пролонгации — оценка" if p["dueEstimated"] else "")],
    ]


def build_workbook(cabinet: dict, simulation: dict | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    cur, nxt = cabinet["currentRate"], cabinet["nextRate"]
    info = cabinet["agentInfo"]

    _sheet(wb, "Сводка", ["Показатель", "Значение"], [
        ["Агент", cabinet["agent"]],
        ["Филиал", f"{info['branch']} ({info['group'] or 'группа не задана'})"],
        ["Данные", f"{cabinet['source']}, факт по {cabinet['asOf']}"],
        ["Правила", cur["rule"]["title"]],
        [f"Ставки {cabinet['quarter']}", cur["statusTitle"]],
        [f"Ставки {cabinet['nextQuarter']}", nxt["statusTitle"]],
        *[[f"АВ {cabinet['quarter']}, {v['scenarioTitle']}", round(v["kv"], 2)] for v in cabinet["income"]["current"].values()],
        *[[f"АВ {cabinet['nextQuarter']}, {v['scenarioTitle']}", round(v["kv"], 2)] for v in cabinet["income"]["next"].values()],
        ["Сценарии", cabinet["scenarioNote"]],
    ])
    header = ["Квартал", "Вид", "База, %", "Доп.АВ1, %", "Доп.АВ2, %", "Доп.АВ3, %", "Итого, %", "Ограничено 30%", "Примечание"]
    _sheet(wb, "Ставки", header, _rates_rows(cur) + _rates_rows(nxt))
    _sheet(wb, "Показатели", ["Ставка квартала", "Показатель", "Значение", "Начислено, %", "Пояснение"],
           _metrics_rows(cur) + _metrics_rows(nxt))
    _sheet(wb, "Доход", ["Квартал", "Сценарий", "Вид", "Сборы", "Ставка, %", "АВ"], [
        [inc["quarter"], inc["scenarioTitle"], g["title"], round(g["base"], 2), _pct(g["rate"]), round(g["kv"], 2)]
        for period in ("current", "next")
        for inc in cabinet["income"][period].values()
        for g in inc["lines"]
    ])

    plan = cabinet["plan"]
    rows = []
    if not plan.get("available", True):
        rows.append(["—", plan["reason"], None, None, None, None])
    for st in plan.get("steps", []):
        n = st["needs"]
        need = (f"+{n['dvs']:,.0f} ₽ ДВС" if "dvs" in n else f"+{n['priority']:,.0f} ₽ приоритетных" if "priority" in n
                else f"{n['prolonged']} из {n['due']} договоров" if "prolonged" in n else st.get("blocked", ""))
        rows.append([st["title"], st["target"], need.replace(",", " "), _pct(st["from"]), _pct(st["to"]), None])
    for opt in plan.get("options", []):
        rows.append([opt["title"], f"КАСКО {opt['kasko']:,.0f} ₽, приоритетные {opt['priority']:,.0f} ₽".replace(",", " "),
                     f"пролонгаций {opt['prolonged'] if opt['prolonged'] is not None else '—'}", None, None,
                     round(opt["effects"]["total"], 2)])
    _sheet(wb, "План", ["Шаг / вариант", "Условие", "Нужно", "Было, %", "Станет, %", "Эффект, ₽"], rows)

    if simulation:
        e = simulation["effects"]
        _sheet(wb, "Сценарий «что если»", ["Показатель", "Текущий", "Сценарий"], [
            ["АВ " + simulation["quarter"], round(simulation["baseline"]["incomeCurrent"]["kv"], 2),
             round(simulation["scenario"]["incomeCurrent"]["kv"], 2)],
            ["АВ " + simulation["nextQuarter"], round(simulation["baseline"]["incomeNext"]["kv"], 2),
             round(simulation["scenario"]["incomeNext"]["kv"], 2)],
            *[[f"Ставка {b['title']} {simulation['nextQuarter']}, %", _pct(b["total"]), _pct(a["total"])]
              for b, a in zip(simulation["baseline"]["nextRate"]["rates"], simulation["scenario"]["nextRate"]["rates"])],
            ["Эффект доп. продаж", None, round(e["salesEffect"], 2)],
            ["Эффект ставки", None, round(e["rateEffect"], 2)],
        ])

    _sheet(wb, "Допущения", ["Допущение"], [[a] for a in cabinet["assumptions"]])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
